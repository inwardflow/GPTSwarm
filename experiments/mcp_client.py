#!/usr/bin/env python3
"""
MCP Client Adapter
==================
A lightweight MCP client that communicates with MCP servers via stdio (JSON-RPC 2.0).
Discovers tools from MCP servers and provides a unified interface for calling them.

This module reads mcp.json and manages MCP server processes.
"""

import os
import json
import subprocess
import threading
import time
import sys


class MCPServerProcess:
    """Manages a single MCP server process via stdio."""
    
    def __init__(self, name: str, command: str, args: list, env: dict = None):
        self.name = name
        self.command = command
        self.args = args
        self.env = env or {}
        self.process = None
        self._request_id = 0
        self._lock = threading.Lock()
        self._responses = {}
        self._reader_thread = None
        self._running = False
    
    def start(self):
        """Start the MCP server process."""
        full_env = os.environ.copy()
        full_env.update(self.env)
        
        try:
            self.process = subprocess.Popen(
                [self.command] + self.args,
                stdin=subprocess.PIPE,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                env=full_env,
                cwd=os.path.dirname(os.path.abspath(__file__))
            )
            self._running = True
            
            # Start reader thread
            self._reader_thread = threading.Thread(target=self._read_responses, daemon=True)
            self._reader_thread.start()
            
            # Initialize the server
            self._initialize()
            return True
        except Exception as e:
            print(f"Failed to start MCP server '{self.name}': {e}", file=sys.stderr)
            return False
    
    def _read_responses(self):
        """Read JSON-RPC responses from the server's stdout."""
        while self._running and self.process and self.process.poll() is None:
            try:
                line = self.process.stdout.readline()
                if not line:
                    break
                line = line.strip()
                if not line:
                    continue
                try:
                    msg = json.loads(line)
                    req_id = msg.get("id")
                    if req_id is not None:
                        self._responses[req_id] = msg
                except json.JSONDecodeError:
                    pass
            except Exception:
                break
    
    def _send_request(self, method: str, params: dict = None, timeout: float = 30) -> dict:
        """Send a JSON-RPC 2.0 request and wait for response."""
        with self._lock:
            self._request_id += 1
            req_id = self._request_id
        
        request = {
            "jsonrpc": "2.0",
            "id": req_id,
            "method": method,
        }
        if params:
            request["params"] = params
        
        request_line = json.dumps(request) + "\n"
        
        try:
            self.process.stdin.write(request_line.encode())
            self.process.stdin.flush()
        except Exception as e:
            return {"error": f"Failed to send request: {e}"}
        
        # Wait for response
        start = time.time()
        while time.time() - start < timeout:
            if req_id in self._responses:
                return self._responses.pop(req_id)
            time.sleep(0.05)
        
        return {"error": f"Timeout waiting for response to {method}"}
    
    def _initialize(self):
        """Send the initialize request to the MCP server."""
        resp = self._send_request("initialize", {
            "protocolVersion": "2024-11-05",
            "capabilities": {},
            "clientInfo": {"name": "gaia-experiment", "version": "1.0.0"}
        })
        
        # Send initialized notification
        notif = json.dumps({"jsonrpc": "2.0", "method": "notifications/initialized"}) + "\n"
        try:
            self.process.stdin.write(notif.encode())
            self.process.stdin.flush()
        except Exception:
            pass
        
        return resp
    
    def list_tools(self) -> list:
        """List available tools from this server."""
        resp = self._send_request("tools/list", {})
        if "result" in resp:
            return resp["result"].get("tools", [])
        return []
    
    def call_tool(self, tool_name: str, arguments: dict, timeout: float = 60) -> str:
        """Call a tool on this server."""
        resp = self._send_request("tools/call", {
            "name": tool_name,
            "arguments": arguments
        }, timeout=timeout)
        
        if "result" in resp:
            content = resp["result"].get("content", [])
            texts = []
            for item in content:
                if isinstance(item, dict) and item.get("type") == "text":
                    texts.append(item.get("text", ""))
                elif isinstance(item, str):
                    texts.append(item)
            return "\n".join(texts) if texts else json.dumps(resp["result"])
        elif "error" in resp:
            return f"MCP error: {json.dumps(resp['error'])}"
        return f"Unexpected response: {json.dumps(resp)}"
    
    def stop(self):
        """Stop the MCP server process."""
        self._running = False
        if self.process:
            try:
                self.process.terminate()
                self.process.wait(timeout=5)
            except Exception:
                try:
                    self.process.kill()
                except Exception:
                    pass


class MCPToolManager:
    """Manages multiple MCP servers and provides unified tool access."""
    
    def __init__(self, config_path: str = None):
        self.servers = {}  # name -> MCPServerProcess
        self.tool_map = {}  # tool_name -> server_name
        self.tool_schemas = {}  # tool_name -> schema dict
        
        if config_path:
            self.load_config(config_path)
    
    def load_config(self, config_path: str):
        """Load MCP server configuration from mcp.json."""
        with open(config_path) as f:
            config = json.load(f)
        
        mcp_servers = config.get("mcpServers", {})
        for name, server_config in mcp_servers.items():
            command = server_config.get("command", "")
            args = server_config.get("args", [])
            env = server_config.get("env", {})
            self.servers[name] = MCPServerProcess(name, command, args, env)
    
    def start_all(self):
        """Start all configured MCP servers and discover tools."""
        for name, server in self.servers.items():
            print(f"Starting MCP server: {name}...", flush=True)
            if server.start():
                time.sleep(1)  # Give server time to initialize
                tools = server.list_tools()
                for tool in tools:
                    tool_name = tool.get("name", "")
                    self.tool_map[tool_name] = name
                    self.tool_schemas[tool_name] = tool
                    print(f"  Discovered tool: {tool_name}", flush=True)
            else:
                print(f"  Failed to start server: {name}", flush=True)
    
    def start_server(self, name: str):
        """Start a specific MCP server."""
        if name in self.servers:
            server = self.servers[name]
            if server.start():
                time.sleep(1)
                tools = server.list_tools()
                for tool in tools:
                    tool_name = tool.get("name", "")
                    self.tool_map[tool_name] = name
                    self.tool_schemas[tool_name] = tool
    
    def get_tool_definitions(self) -> list:
        """Get OpenAI-compatible tool definitions for all discovered tools."""
        definitions = []
        for tool_name, schema in self.tool_schemas.items():
            definitions.append({
                "type": "function",
                "function": {
                    "name": schema.get("name", tool_name),
                    "description": schema.get("description", ""),
                    "parameters": schema.get("inputSchema", {
                        "type": "object",
                        "properties": {},
                        "required": []
                    })
                }
            })
        return definitions
    
    def call_tool(self, tool_name: str, arguments: dict, timeout: float = 60) -> str:
        """Call a tool by name, routing to the correct server."""
        server_name = self.tool_map.get(tool_name)
        if not server_name:
            return f"Unknown tool: {tool_name}"
        
        server = self.servers.get(server_name)
        if not server:
            return f"Server not found: {server_name}"
        
        return server.call_tool(tool_name, arguments, timeout=timeout)
    
    def stop_all(self):
        """Stop all MCP servers."""
        for name, server in self.servers.items():
            server.stop()


class DirectToolManager:
    """
    A direct (in-process) tool manager that provides the same interface as MCPToolManager
    but calls tools directly without going through MCP stdio protocol.
    
    This is more reliable and faster, while maintaining the same mcp.json configuration
    for documentation and reproducibility.
    """
    
    def __init__(self):
        self.tools = {}
        self._setup_tools()
    
    def _setup_tools(self):
        """Register all tools with their implementations."""
        self.tools = {
            "search_web": {
                "name": "search_web",
                "description": "Search the web using DuckDuckGo. Returns titles, URLs, and snippets for the top results.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "query": {"type": "string", "description": "The search query string"},
                        "max_results": {"type": "integer", "description": "Maximum results (default: 5)", "default": 5}
                    },
                    "required": ["query"]
                },
                "fn": self._search_web
            },
            "read_file": {
                "name": "read_file",
                "description": "Read and extract content from various file types. Supports: PDF, Excel (.xlsx/.xls), CSV, DOCX, PPTX, ZIP, JSON, TXT, PY, MP3/WAV (audio transcription), PDB, JSONLD, and other text formats.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "file_path": {"type": "string", "description": "The absolute path to the file to read"}
                    },
                    "required": ["file_path"]
                },
                "fn": self._read_file
            },
            "run_python": {
                "name": "run_python",
                "description": "Execute Python code and return stdout/stderr. Use for calculations, data processing, or any computation. Available packages: numpy, pandas, openpyxl, matplotlib, scipy, etc.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "code": {"type": "string", "description": "Python code to execute"}
                    },
                    "required": ["code"]
                },
                "fn": self._run_python
            },
            "analyze_image": {
                "name": "analyze_image",
                "description": "Analyze an image file using OCR (Optical Character Recognition). Extracts text content from images.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "file_path": {"type": "string", "description": "The absolute path to the image file (PNG, JPG, etc.)"}
                    },
                    "required": ["file_path"]
                },
                "fn": self._analyze_image
            },
            "fetch_webpage": {
                "name": "fetch_webpage",
                "description": "Fetch and extract text content from a webpage URL. Useful for reading articles, documentation, or any web page content.",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "url": {"type": "string", "description": "The URL of the webpage to fetch"}
                    },
                    "required": ["url"]
                },
                "fn": self._fetch_webpage
            }
        }
    
    def get_tool_definitions(self) -> list:
        """Get OpenAI-compatible tool definitions."""
        definitions = []
        for name, tool in self.tools.items():
            definitions.append({
                "type": "function",
                "function": {
                    "name": tool["name"],
                    "description": tool["description"],
                    "parameters": tool["parameters"]
                }
            })
        return definitions
    
    def call_tool(self, tool_name: str, arguments: dict, timeout: float = 60) -> str:
        """Call a tool by name."""
        tool = self.tools.get(tool_name)
        if not tool:
            return f"Unknown tool: {tool_name}"
        try:
            return tool["fn"](arguments)
        except Exception as e:
            return f"Tool error: {e}"
    
    # ---- Tool Implementations ----
    
    def _search_web(self, args: dict) -> str:
        query = args.get("query", "")
        max_results = args.get("max_results", 5)
        try:
            try:
                from ddgs import DDGS
            except ImportError:
                from duckduckgo_search import DDGS
            ddgs = DDGS()
            results = list(ddgs.text(query, max_results=max_results))
            if not results:
                return "No search results found."
            output = []
            for i, r in enumerate(results, 1):
                output.append(f"[{i}] {r.get('title', 'N/A')}")
                output.append(f"    URL: {r.get('href', 'N/A')}")
                output.append(f"    {r.get('body', 'N/A')}")
            return "\n".join(output)
        except Exception as e:
            return f"Search error: {e}"
    
    def _read_file(self, args: dict) -> str:
        file_path = args.get("file_path", "")
        if not os.path.exists(file_path):
            return f"File not found: {file_path}"
        
        ext = file_path.rsplit(".", 1)[-1].lower() if "." in file_path else ""
        
        try:
            if ext == "pdf":
                return self._read_pdf(file_path)
            elif ext in ("xlsx", "xls"):
                return self._read_excel(file_path)
            elif ext == "csv":
                with open(file_path, "r", errors="replace") as f:
                    return f.read()[:10000]
            elif ext == "docx":
                try:
                    from docx import Document
                    doc = Document(file_path)
                    return "\n".join(p.text for p in doc.paragraphs)[:10000]
                except ImportError:
                    return "python-docx not installed."
            elif ext == "pptx":
                try:
                    from pptx import Presentation
                    prs = Presentation(file_path)
                    text = []
                    for slide in prs.slides:
                        for shape in slide.shapes:
                            if hasattr(shape, "text"):
                                text.append(shape.text)
                    return "\n".join(text)[:10000]
                except ImportError:
                    return "python-pptx not installed."
            elif ext == "zip":
                import zipfile
                with zipfile.ZipFile(file_path, "r") as zf:
                    names = zf.namelist()
                    output = [f"ZIP contents ({len(names)} files):"]
                    for name in names[:20]:
                        output.append(f"  {name}")
                    for name in names:
                        if name.endswith((".txt", ".csv", ".json", ".py", ".md")):
                            try:
                                content = zf.read(name).decode("utf-8", errors="replace")
                                output.append(f"\n--- {name} ---")
                                output.append(content[:3000])
                            except Exception:
                                pass
                    return "\n".join(output)[:10000]
            elif ext in ("mp3", "wav", "mp4", "webm"):
                try:
                    result = subprocess.run(
                        ["manus-speech-to-text", file_path],
                        capture_output=True, text=True, timeout=120
                    )
                    if result.stdout.strip():
                        return result.stdout.strip()[:5000]
                    return f"Audio transcription returned empty. stderr: {result.stderr[:500]}"
                except Exception as e:
                    return f"Audio transcription error: {e}"
            elif ext in ("png", "jpg", "jpeg", "gif", "bmp", "tiff"):
                return self._do_ocr(file_path)
            else:
                with open(file_path, "r", errors="replace") as f:
                    return f.read()[:10000]
        except Exception as e:
            return f"Error reading file: {e}"
    
    def _read_pdf(self, file_path: str) -> str:
        try:
            result = subprocess.run(
                ["pdftotext", "-layout", file_path, "-"],
                capture_output=True, text=True, timeout=30
            )
            text = result.stdout.strip()
            if text:
                return text[:10000]
        except Exception:
            pass
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(file_path)
            text = ""
            for page in reader.pages:
                text += page.extract_text() or ""
            return text[:10000] if text.strip() else "PDF appears to contain only images."
        except Exception as e:
            return f"PDF read error: {e}"
    
    def _read_excel(self, file_path: str) -> str:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        output = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            output.append(f"=== Sheet: {sheet_name} ===")
            for row in ws.iter_rows(values_only=True):
                output.append("\t".join(str(c) if c is not None else "" for c in row))
        return "\n".join(output)[:15000]
    
    def _do_ocr(self, file_path: str) -> str:
        try:
            import pytesseract
            from PIL import Image
            img = Image.open(file_path)
            info = f"Image: {os.path.basename(file_path)}, Size: {img.size}, Mode: {img.mode}"
            text = pytesseract.image_to_string(img)
            if text.strip():
                return f"{info}\nOCR Text:\n{text.strip()}"
            return f"{info}\nOCR returned no text."
        except ImportError:
            from PIL import Image
            img = Image.open(file_path)
            return f"Image: {os.path.basename(file_path)}, Size: {img.size}, Mode: {img.mode}. (OCR not available)"
        except Exception as e:
            return f"OCR error: {e}"
    
    def _run_python(self, args: dict) -> str:
        code = args.get("code", "")
        try:
            import tempfile
            with tempfile.NamedTemporaryFile(mode="w", suffix=".py", delete=False) as f:
                f.write(code)
                f.flush()
                result = subprocess.run(
                    ["python3", f.name],
                    capture_output=True, text=True, timeout=30, cwd="/tmp"
                )
                os.unlink(f.name)
                output = result.stdout
                if result.stderr:
                    output += "\nSTDERR: " + result.stderr
                return output.strip()[:5000] if output.strip() else "(No output)"
        except subprocess.TimeoutExpired:
            return "Code execution timed out (30s limit)."
        except Exception as e:
            return f"Code execution error: {e}"
    
    def _analyze_image(self, args: dict) -> str:
        file_path = args.get("file_path", "")
        if not os.path.exists(file_path):
            return f"File not found: {file_path}"
        return self._do_ocr(file_path)
    
    def _fetch_webpage(self, args: dict) -> str:
        url = args.get("url", "")
        try:
            import httpx
            resp = httpx.get(
                url, timeout=30, verify=False, follow_redirects=True,
                headers={"User-Agent": "Mozilla/5.0 (compatible; research-bot)"}
            )
            if resp.status_code != 200:
                return f"HTTP {resp.status_code} error fetching {url}"
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(resp.text, "html.parser")
            for tag in soup(["script", "style", "nav", "footer", "header"]):
                tag.decompose()
            text = soup.get_text(separator="\n", strip=True)
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            return "\n".join(lines)[:8000]
        except Exception as e:
            return f"Webpage fetch error: {e}"
    
    def stop_all(self):
        """No-op for direct tool manager."""
        pass
